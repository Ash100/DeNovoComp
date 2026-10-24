# -*- coding: utf-8 -*-
"""Build the cover letter for the Nature Computational Science submission."""
import json
from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

import manuscript_prose as P

ROOT = Path("/home/claude/dnc")
OUT = ROOT / "out"
DELIV = ROOT / "deliverables"; DELIV.mkdir(exist_ok=True)
R = json.loads((OUT / "run_record.json").read_text())

CAMP = ["A", "B", "C", "D"]
cam, feas, tri = R["campaign"], R["feasibility"], R["triage"]
bench, scaf = R["benchmarks"], R["scaffold"]
cur, pri, cfg = R["curation"], R["prior"], R["config"]


def pct(x, d=0): return f"{100 * float(x):.{d}f}%"
def fold(a, b): return f"{b / a:.1f}-fold" if a > 1e-9 else "from a near-zero baseline"


def setup(doc):
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(11.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    st.paragraph_format.line_spacing = 1.15
    st.paragraph_format.space_after = Pt(9)
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Inches(1.0)
    sec.top_margin = sec.bottom_margin = Inches(0.9)


def para(doc, text, *, bold=False, italic=False, size=11.5, after=9, align=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold; r.italic = italic; r.font.size = Pt(size)
    r.font.name = "Times New Roman"
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.15
    if align is not None:
        p.alignment = align
    return p


def J(doc, text, **kw):
    return para(doc, text, align=WD_ALIGN_PARAGRAPH.JUSTIFY, **kw)


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(text); r.font.size = Pt(11.5); r.font.name = "Times New Roman"
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(7)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p


def build():
    doc = Document()
    setup(doc)

    para(doc, "Department of Bioinformatics", size=10.5, after=0)
    para(doc, "Hazara University, Mansehra", size=10.5, after=0)
    para(doc, "Khyber Pakhtunkhwa, Pakistan", size=10.5, after=0)
    para(doc, "ashfaq333@gmail.com", size=10.5, after=14)
    para(doc, "24 October 2026", after=14)
    para(doc, "The Editor", after=0)
    para(doc, "Nature Computational Science", after=14)
    para(doc, "Dear Editor,", after=12)

    J(doc, "I am pleased to submit the enclosed manuscript, “" + P.TITLE +
           "”, for consideration as an Article in Nature Computational Science.")

    J(doc, "Goal-directed generative models now reliably produce molecules that satisfy a "
           "numerical objective. The step that follows, deciding which of those molecules "
           "a project should pursue and being able to justify that decision, is usually "
           "implemented as project-specific code that is not published. The thresholds, "
           "weights and filter orderings that determine a shortlist are consequently "
           "difficult for others to reproduce or to evaluate.")

    J(doc, "The manuscript describes DeNovoComp, a pipeline in which design intent is "
           "declared rather than implemented. A single specification of property windows, "
           "similarity bounds and substructural constraints is compiled, through "
           "Derringer–Suich desirability functions, into an object that serves "
           "simultaneously as the reinforcement-learning reward, the ranking function and "
           "the provenance record of the campaign. Four features of the work seem to us "
           "worth the journal's attention:")

    for txt in [
        "A feasibility assessment carried out before optimisation. The unsteered prior is "
        "scored against the specification before any computational budget is committed, "
        "and shortfalls are attributed to individual objectives. Reinforcement learning "
        "against an unreachable objective produces no error condition, only a flat reward "
        "trajectory and a library indistinguishable from the baseline, which is readily "
        "misread as a limitation of the method rather than of the request. In our "
        f"campaigns the preview distinguished a demanding specification from an infeasible "
        f"one: specification D returned a baseline median of {feas['D']['median']:.3f}, "
        f"because only {pct(feas['D']['gate_pass_fraction'])} of molecules contain the "
        f"mandatory substructure, yet its 95th percentile of {feas['D']['q95']:.3f} "
        f"indicated an attainable target, which optimisation then reached.",

        "Synthesizability as an optimisation objective rather than a downstream filter, "
        "following the demonstration by Gao and Coley that generative models frequently "
        "propose molecules that cannot readily be prepared and that standard benchmarks "
        "do not detect this. Leverage-based applicability-domain membership is likewise "
        "computed for every generated molecule and propagated into the triage rule.",

        "A measured compliance-versus-diversity frontier. The same instrumentation that "
        "records gains in compliance also showed that the default reinforcement-learning "
        "intensity over-optimises: in three of the four campaigns uniqueness fell to "
        f"between {min(bench['designed_' + n]['uniqueness'] for n in CAMP):.2f} and "
        f"{sorted(bench['designed_' + n]['uniqueness'] for n in CAMP)[2]:.2f}, distinct "
        f"Bemis–Murcko frameworks fell from "
        f"{int(scaf['baseline']['n_scaffolds']):,} to as few as "
        f"{int(min(scaf['designed_' + n]['n_scaffolds'] for n in CAMP))}, "
        "structural-alert incidence increased, and the unsteered baseline yielded more "
        "GO-labelled molecules than any steered library. We therefore mapped the frontier "
        "by taking policy snapshots at four intensities, and report settings that retain "
        "most of the attainable compliance at close to baseline diversity.",

        "Reporting of results that run counter to the headline measure. Molecular-weight "
        "compliance decreased in every campaign, falling below the requested window, even "
        "as the composite design score rose, because that is what the aggregation rule "
        "traded away. Reported by composite score alone, these four campaigns would "
        "appear uniformly successful.",
    ]:
        bullet(doc, txt)

    J(doc, f"We report four specifications declared in advance and executed against a "
           f"shared prior and a shared unsteered control library: general drug-like, "
           f"CNS-oriented, analogue design against a reference ligand, and "
           f"substructure-constrained. REINFORCE fine-tuning increased the mean composite "
           f"design score in every campaign, by "
           f"{fold(cam['A']['baseline_mean'], cam['A']['designed_mean'])} in the general "
           f"case, and the mandatory-substructure gate raised compliance from "
           f"{pct(cam['D']['baseline_gate_pass'])} to "
           f"{pct(cam['D']['designed_gate_pass'])} without any post-generation filtering.")

    J(doc, "We believe the work suits the journal's readership for two reasons. The "
           "problem it addresses, that the limiting step in a maturing computational field "
           "has moved from model capability to evaluation, prioritisation and "
           "auditability, is not specific to chemistry, and the pattern of declaring "
           "intent as a specification that doubles as a provenance record generalises to "
           "other domains in which the output of a generative model must be defended to a "
           "decision-maker. The pipeline is also executable on freely available hosted "
           "infrastructure, with documented degradation when optional components are "
           "unavailable, which makes it usable for teaching as well as research.")

    J(doc, f"We have been explicit about what the work does not establish. The objective "
           f"set is physicochemical and structural and contains no docking or activity "
           f"model, and a GO label is defined narrowly as computational priority under a "
           f"stated rule set. The reported run was executed on CPU-only hardware, which "
           f"required the training corpus to be subsampled to {cur['final_training']:,} "
           f"of {cur['after_dedup']:,} curated molecules; the full configuration is "
           f"tabulated so that the scale of the results is unambiguous. SCScore and "
           f"AiZynthFinder could not be installed in that environment, and the manuscript "
           f"states this rather than omitting it. The leverage-based "
           f"applicability-domain test did not discriminate at this configuration, since "
           f"all molecules fell within the threshold, and we report that rather than "
           f"presenting the check as satisfied. No prospective synthesis or assay has been "
           f"performed, and we identify prospective validation as the decisive untested "
           f"question.")

    J(doc, "The manuscript is original, has not been published elsewhere and is not under "
           "consideration by another journal. [Please confirm or amend: no related "
           "manuscripts by any author are under consideration or in press elsewhere, and "
           "there have been no prior discussions with a Nature Computational Science "
           "editor regarding this work.] All authors have approved the submission and "
           "declare no competing interests. Use of a large language model as a coding and "
           "drafting assistant is documented in the Methods; it is not an author, and no "
           "reported result was produced by it.")

    J(doc, "[Optional: suggested reviewers, with affiliations, email addresses and a "
           "one-line justification for each; and any requested exclusions, with reasons.]",
      italic=True)

    J(doc, "Thank you for considering this submission.", after=18)

    para(doc, "Yours sincerely,", after=28)
    para(doc, "Ashfaq Ahmad, PhD", bold=True, after=0)
    para(doc, "Assistant Professor of Bioinformatics", size=10.5, after=0)
    para(doc, "Department of Bioinformatics, Hazara University, Mansehra, Pakistan",
         size=10.5, after=0)
    para(doc, "On behalf of all authors", size=10.5, italic=True, after=0)

    out = DELIV / "DeNovoComp_Cover_Letter_NatComputSci.docx"
    doc.save(out)
    return out


if __name__ == "__main__":
    print("WROTE", build())
